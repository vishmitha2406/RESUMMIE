from flask import Flask, request, jsonify, render_template, send_from_directory
import os
import re
import json
from datetime import datetime
from PyPDF2 import PdfReader
from docx import Document
from werkzeug.utils import secure_filename

from models import db, Resume

# ============================================================
# APP CONFIGURATION
# ============================================================
app = Flask(__name__)

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB
app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{os.path.join(BASE_DIR, "resumes.db")}'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

ALLOWED_EXTENSIONS = {'pdf', 'docx'}

# Initialize database
db.init_app(app)

with app.app_context():
    db.create_all()
    print("✅ Database initialized at:", app.config['SQLALCHEMY_DATABASE_URI'])


# ============================================================
# JOB ROLES DATABASE
# ============================================================
JOB_ROLES = {
    "Software Engineer": {
        "skills": ["python", "java", "c++", "data structures", "algorithms",
                   "git", "sql", "oop", "problem solving"],
        "extra_skills": ["system design", "microservices", "cloud computing",
                         "ci/cd", "testing", "docker"],
        "description": "Develops and maintains software applications with clean, efficient code."
    },
    "Web Developer": {
        "skills": ["html", "css", "javascript", "react", "angular", "vue",
                   "node", "express", "bootstrap"],
        "extra_skills": ["typescript", "webpack", "graphql", "responsive design",
                         "seo", "accessibility"],
        "description": "Builds and maintains websites with focus on user experience."
    },
    "Data Scientist": {
        "skills": ["python", "r", "sql", "machine learning", "deep learning",
                   "nlp", "computer vision", "statistics", "pandas", "numpy"],
        "extra_skills": ["big data", "spark", "hadoop", "cloud computing",
                         "data visualization", "tableau"],
        "description": "Analyzes complex data to extract insights and build predictive models."
    },
    "DevOps Engineer": {
        "skills": ["aws", "azure", "gcp", "docker", "kubernetes", "jenkins",
                   "git", "linux", "ci/cd", "ansible"],
        "extra_skills": ["terraform", "prometheus", "grafana", "elk stack",
                         "security", "networking"],
        "description": "Manages infrastructure and automates deployment processes."
    },
    "Data Analyst": {
        "skills": ["sql", "excel", "python", "r", "tableau", "power bi",
                   "statistics", "data visualization", "analytical"],
        "extra_skills": ["big data", "etl", "data warehousing",
                         "business intelligence", "reporting"],
        "description": "Interprets data and creates reports to support business decisions."
    },
    "Frontend Developer": {
        "skills": ["html", "css", "javascript", "react", "angular", "vue",
                   "sass", "less", "bootstrap", "jquery"],
        "extra_skills": ["typescript", "webpack", "babel", "state management",
                         "ui/ux", "responsive design"],
        "description": "Creates user interfaces and ensures smooth interactions."
    },
    "Backend Developer": {
        "skills": ["python", "java", "c#", "node", "express", "django",
                   "flask", "spring", "sql", "mongodb"],
        "extra_skills": ["microservices", "graphql", "redis", "kafka",
                         "rabbitmq", "security"],
        "description": "Builds server-side logic, APIs, and database interactions."
    },
    "Full Stack Developer": {
        "skills": ["html", "css", "javascript", "react", "node", "python",
                   "sql", "mongodb", "git", "rest api"],
        "extra_skills": ["typescript", "docker", "aws", "graphql",
                         "testing", "ci/cd"],
        "description": "Works on both frontend and backend development."
    },
    "Mobile Developer": {
        "skills": ["react native", "flutter", "swift", "kotlin", "java",
                   "android", "ios", "mobile", "firebase"],
        "extra_skills": ["ui/ux design", "app store deployment",
                         "performance optimization", "testing"],
        "description": "Builds mobile applications for iOS and Android."
    },
    "Machine Learning Engineer": {
        "skills": ["python", "tensorflow", "pytorch", "scikit-learn",
                   "machine learning", "deep learning", "nlp", "computer vision"],
        "extra_skills": ["mlops", "docker", "kubernetes", "big data",
                         "cloud computing", "model deployment"],
        "description": "Designs and implements machine learning models."
    },
    "Cloud Engineer": {
        "skills": ["aws", "azure", "gcp", "docker", "kubernetes", "linux",
                   "python", "networking", "security"],
        "extra_skills": ["terraform", "ansible", "prometheus",
                         "cloud architecture", "cost optimization"],
        "description": "Designs and manages cloud infrastructure."
    },
    "Business Analyst": {
        "skills": ["excel", "sql", "python", "tableau", "business analysis",
                   "requirements gathering", "documentation"],
        "extra_skills": ["agile", "scrum", "project management",
                         "communication", "stakeholder management"],
        "description": "Analyzes business needs and defines requirements."
    },
    "Project Manager": {
        "skills": ["agile", "scrum", "project management", "leadership",
                   "communication", "jira", "confluence"],
        "extra_skills": ["budgeting", "risk management",
                         "stakeholder management", "reporting", "scheduling"],
        "description": "Leads projects and coordinates teams."
    }
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================
def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_pdf(file_path):
    """Extract text from a PDF file."""
    text = ""
    try:
        with open(file_path, 'rb') as f:
            reader = PdfReader(f)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
    except Exception as e:
        print(f"❌ PDF read error: {e}")
    return text


def extract_text_from_docx(file_path):
    """Extract text from a DOCX file."""
    text = ""
    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"❌ DOCX read error: {e}")
    return text


def extract_email(text):
    pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    matches = re.findall(pattern, text)
    return matches[0] if matches else "Not found"


def extract_phone(text):
    patterns = [
        r'\+\d{1,3}[-.\s]?\(?\d{1,4}\)?[-.\s]?\d{1,4}[-.\s]?\d{1,9}',
        r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',
        r'\d{3}[-.\s]\d{3}[-.\s]\d{4}',
        r'\b\d{10,15}\b'
    ]
    for pattern in patterns:
        matches = re.findall(pattern, text)
        if matches:
            phone = re.sub(r'\s+', ' ', matches[0].strip())
            return phone
    return "Not found"


def extract_name(text):
    """Try to find candidate name."""
    lines = [l.strip() for l in text.split('\n') if l.strip()]

    for line in lines[:15]:
        m = re.search(r'(?:Name|Candidate|Full Name)\s*[:\-]\s*([A-Za-z][A-Za-z\s\.]{2,40})',
                      line, re.IGNORECASE)
        if m:
            return m.group(1).strip()

    for line in lines[:10]:
        if len(line) > 50:
            continue
        words = line.split()
        if 2 <= len(words) <= 4 and all(w[0].isupper() for w in words if w):
            if not re.search(r'\d|@|http', line):
                return line
    return "Not found"


def extract_skills(text):
    skill_keywords = [
        'python', 'java', 'javascript', 'react', 'angular', 'vue', 'node',
        'django', 'flask', 'spring', 'aws', 'azure', 'gcp', 'docker',
        'kubernetes', 'sql', 'mongodb', 'postgresql', 'mysql', 'redis',
        'html', 'css', 'typescript', 'jquery', 'bootstrap',
        'c++', 'c#', 'ruby', 'php', 'swift', 'kotlin', 'rust', 'go',
        'machine learning', 'deep learning', 'nlp', 'computer vision',
        'tensorflow', 'pytorch', 'scikit-learn', 'pandas', 'numpy',
        'agile', 'scrum', 'jira', 'git', 'github',
        'jenkins', 'linux', 'unix',
        'project management', 'leadership', 'communication',
        'problem solving', 'analytical', 'teamwork',
        'excel', 'powerpoint', 'word', 'outlook',
        'salesforce', 'sap', 'oracle', 'tableau', 'power bi',
        'html5', 'css3', 'sass', 'less', 'webpack', 'babel',
        'rest api', 'graphql', 'microservices', 'devops',
        'c', 'scala', 'perl', 'matlab',
        'hadoop', 'spark', 'kafka', 'rabbitmq', 'nginx',
        'react native', 'flutter', 'android', 'ios', 'firebase',
        'ui/ux', 'figma', 'photoshop', 'illustrator',
        'data analysis', 'statistics', 'data science',
        'etl', 'data warehousing', 'big data',
        'cloud', 'networking', 'security', 'cryptography'
    ]
    text_lower = text.lower()
    found = []
    for skill in skill_keywords:
        if re.search(r'\b' + re.escape(skill) + r'\b', text_lower):
            skill_title = ' '.join(w.capitalize() for w in skill.split())
            if skill_title not in found:
                found.append(skill_title)
    return found


def extract_experience(text):
    patterns = [
        r'(\d+)\s*\+?\s*(?:years?|yrs?)\s*(?:of\s*)?experience',
        r'experience\s*(?:of\s*)?(\d+)\s*\+?\s*(?:years?|yrs?)',
        r'(\d+)\s*\+?\s*(?:years?|yrs?)\s*experience'
    ]
    for pattern in patterns:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            return f"{m.group(1)} years"

    years = sorted(set(int(y) for y in re.findall(r'(?:19|20)\d{2}', text)))
    if len(years) >= 2 and years[-1] > years[0]:
        return f"{years[-1] - years[0]} years"
    return "Not specified"


def analyze_job_role(skills):
    """Score resume skills against each job role."""
    if not skills:
        return []

    skills_lower = [s.lower() for s in skills]
    results = []

    for role_name, role_data in JOB_ROLES.items():
        required = [s.lower() for s in role_data['skills']]
        matched = [s for s in required
                   if any(s in sk or sk in s for sk in skills_lower)]

        match_pct = round((len(matched) / len(required)) * 100, 1) if required else 0

        missing = [s.title() for s in required
                   if not any(s in sk or sk in s for sk in skills_lower)]

        recommended = [s.title() for s in role_data.get('extra_skills', [])
                       if not any(s.lower() in sk or sk in s.lower()
                                  for sk in skills_lower)]

        results.append({
            'role': role_name,
            'match_percentage': match_pct,
            'matched_skills': [s.title() for s in matched],
            'missing_skills': missing,
            'recommended_skills': recommended[:5],
            'description': role_data['description'],
            'total_required': len(required)
        })

    results.sort(key=lambda x: x['match_percentage'], reverse=True)
    return results[:5]


# ============================================================
# ROUTES
# ============================================================
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/analyze', methods=['POST'])
def analyze_resume():
    """Handle file upload, extract data, save to DB."""
    if 'resume' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400

    file = request.files['resume']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    if not allowed_file(file.filename):
        return jsonify({'error': 'Only PDF and DOCX files are supported'}), 400

    original_name = file.filename
    safe_name = secure_filename(original_name)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    saved_name = f"{timestamp}_{safe_name}"
    saved_path = os.path.join(app.config['UPLOAD_FOLDER'], saved_name)

    try:
        file.save(saved_path)
        file_size = os.path.getsize(saved_path)
    except Exception as e:
        return jsonify({'error': f'Could not save file: {e}'}), 500

    if saved_name.lower().endswith('.pdf'):
        text = extract_text_from_pdf(saved_path)
        file_type = 'pdf'
    else:
        text = extract_text_from_docx(saved_path)
        file_type = 'docx'

    if not text or len(text.strip()) < 50:
        return jsonify({'error': 'Could not extract text from resume.'}), 400

    name = extract_name(text)
    email = extract_email(text)
    contact = extract_phone(text)
    skills = extract_skills(text)
    experience = extract_experience(text)
    role_analysis = analyze_job_role(skills)

    top_role = role_analysis[0]['role'] if role_analysis else None
    top_match = role_analysis[0]['match_percentage'] if role_analysis else None

    try:
        record = Resume(
            filename=original_name,
            file_path=saved_path,
            file_size=file_size,
            file_type=file_type,
            name=name,
            email=email,
            contact=contact,
            experience=experience,
            top_role=top_role,
            top_role_match=top_match,
            raw_text=text[:20000],
            ip_address=request.remote_addr
        )
        record.skills = skills
        record.role_analysis = role_analysis

        db.session.add(record)
        db.session.commit()
        resume_id = record.id
        print(f"💾 Saved resume #{resume_id} ({name})")
    except Exception as e:
        db.session.rollback()
        print(f"❌ DB save error: {e}")
        resume_id = None

    return jsonify({
        'id': resume_id,
        'name': name,
        'email': email,
        'contact': contact,
        'skills': skills or ["No skills found"],
        'experience': experience,
        'role_analysis': role_analysis,
        'top_role': role_analysis[0] if role_analysis else None
    })


@app.route('/history', methods=['GET'])
def history():
    """List all previously analyzed resumes."""
    try:
        records = Resume.query.order_by(Resume.uploaded_at.desc()).limit(50).all()
        return jsonify([r.to_dict() for r in records])
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/resume/<int:resume_id>', methods=['GET'])
def get_resume(resume_id):
    """Get a single saved resume."""
    record = Resume.query.get(resume_id)
    if not record:
        return jsonify({'error': 'Resume not found'}), 404
    return jsonify(record.to_dict(include_text=True))


@app.route('/resume/<int:resume_id>', methods=['DELETE'])
def delete_resume(resume_id):
    """Delete a saved resume."""
    record = Resume.query.get(resume_id)
    if not record:
        return jsonify({'error': 'Resume not found'}), 404

    if record.file_path and os.path.exists(record.file_path):
        try:
            os.remove(record.file_path)
        except Exception as e:
            print(f"⚠️ Could not delete file: {e}")

    db.session.delete(record)
    db.session.commit()
    return jsonify({'success': True, 'message': f'Resume #{resume_id} deleted'})


@app.route('/download/<int:resume_id>', methods=['GET'])
def download_resume(resume_id):
    """Download the original file."""
    record = Resume.query.get(resume_id)
    if not record or not record.file_path or not os.path.exists(record.file_path):
        return jsonify({'error': 'File not found'}), 404
    directory = os.path.dirname(record.file_path)
    filename = os.path.basename(record.file_path)
    return send_from_directory(directory, filename, as_attachment=True)


# ============================================================
# RUN
# ============================================================
if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)